"""Fluxo local: nomes dos arquivos, planilha, PDF e coleta deste computador."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from hw_inventory.ad import domain_from_realm, domain_from_sssd, user_is_domain_admin
from hw_inventory.collect import (
    blank_serial,
    collect_hardware,
    domain_group_label,
    edid_descriptor,
    edid_manufacture,
    edid_native_resolution,
    edid_numeric_serial,
    edid_product_code,
    os_edition_version,
    cpu_place_and_socket,
    parse_dmidecode,
    parse_xrandr_monitors,
    chipset_from_name,
    gpu_instalacao,
    slot_is_free,
    slot_purpose,
    ap_mode_supported,
    hdmi_audio_active,
    nvme_link_label,
    sata_port_is_real,
    sata_speed_from_caps,
    unallocated_from_partitions,
    parse_xrandr_refresh,
    saidas_from_codec,
    usb_has_bluetooth,
    volume_summary,
)
from hw_inventory.models import FormData, Hardware
from hw_inventory.paths import output_filenames, safe_token
from hw_inventory.license_key import (
    is_free_license,
    key_from_msdm,
    license_from_copyright,
    partial_key,
    token_from_ubuntu_pro,
)
from hw_inventory.report import (
    _radio_rows,
    needs_patrimonio_confirmation,
    validate_form,
    write_inventory,
)

_DMI = """
Handle 0x0000, DMI type 0, 26 bytes
BIOS Information
	Vendor: Login Informatica LTDA
	Version: 5.35
	Release Date: 01/02/2024

Handle 0x0001, DMI type 1, 27 bytes
System Information
	Manufacturer: Login Informatica LTDA
	Product Name: LOG-A620-DDR5
	Serial Number: To Be Filled By O.E.M.
	UUID: Not Specified

Handle 0x0002, DMI type 2, 15 bytes
Baseboard Information
	Manufacturer: Login Informatica LTDA
	Product Name: LOG-A620-DDR5
	Serial Number: PLACA123

Handle 0x0010, DMI type 16, 23 bytes
Physical Memory Array
	Number Of Devices: 4
	Maximum Capacity: 128 GB

Handle 0x0011, DMI type 17, 92 bytes
Memory Device
	Size: 16 GB
	Locator: DIMM_A1
	Type: DDR5
	Speed: 5200 MT/s
	Manufacturer: Kingston
	Serial Number: MEM123
	Part Number: KF556

Handle 0x0012, DMI type 17, 92 bytes
Memory Device
	Size: No Module Installed
	Locator: DIMM_A2
	Type: Unknown
	Serial Number: Not Specified

Handle 0x0004, DMI type 4, 48 bytes
Processor Information
	Status: Populated, Enabled
	Manufacturer: AMD
	Serial Number: Not Specified
	Core Count: 8
	Thread Count: 16
"""


def _hardware() -> Hardware:
    return Hardware(
        hostname="pc-lab",
        os_name="Linux Mint 22.2",
        kernel="6.8.0",
        arch="x86_64",
        manufacturer="Login",
        product="LOG-A620",
        version="1.0",
        serial="SERIE-PC",
        uuid="uuid-1",
        bios_vendor="Login",
        bios_version="5.35",
        bios_date="01/02/2024",
        board_vendor="Login",
        board_name="LOG-A620",
        board_version="1.0",
        board_serial="",
        chassis_type="Desktop",
        chassis_serial="",
        cpu_model="AMD Ryzen 7 8700G",
        cpu_vendor="AMD",
        cpu_cores="8",
        cpu_threads="16",
        cpu_max_mhz="5177 MHz",
        cpu_serial="",
        memory_total="16 GB",
        memory_slots_total="4",
        memory_slots_used="1",
        memory_max="128 GB",
        memory_expansion="1 de 4 slots em uso; 3 slot(s) livre(s)",
        collected_at="14/10/2026 10:00",
    )


class NamesTest(unittest.TestCase):
    def test_com_tombamento(self) -> None:
        spreadsheet, pdf = output_filenames("123456", "pc-lab")
        self.assertEqual(spreadsheet, "123456_PC-LAB_inventario.xlsx")
        self.assertEqual(pdf, "123456_PC-LAB_relatorio_hardware.pdf")

    def test_sem_tombamento(self) -> None:
        spreadsheet, pdf = output_filenames("  ", "Meu PC")
        self.assertEqual(spreadsheet, "MEU-PC_inventario.xlsx")
        self.assertEqual(pdf, "MEU-PC_relatorio_hardware.pdf")

    def test_remove_barras(self) -> None:
        self.assertEqual(safe_token("12/345"), "12-345")


class FormTest(unittest.TestCase):
    def test_exige_responsavel_e_local(self) -> None:
        self.assertIsNotNone(validate_form(FormData(patrimonio="1")))
        form = FormData(responsavel="Maria Silva", local="Sala 12")
        self.assertIsNone(validate_form(form))
        self.assertTrue(needs_patrimonio_confirmation(form))
        self.assertFalse(needs_patrimonio_confirmation(FormData(patrimonio="10", responsavel="A", local="B")))


class DmiTest(unittest.TestCase):
    def test_parser_e_serie_vazia(self) -> None:
        records = parse_dmidecode(_DMI)
        types = [item["type"] for item in records]
        self.assertIn(17, types)
        self.assertEqual(blank_serial("To Be Filled By O.E.M."), "")
        self.assertEqual(blank_serial("PLACA123"), "PLACA123")
        modules = [item for item in records if item["type"] == 17]
        self.assertEqual(modules[0]["fields"]["Locator"], "DIMM_A1")
        self.assertEqual(modules[1]["fields"]["Size"], "No Module Installed")


class AdParseTest(unittest.TestCase):
    def test_dominio_e_grupo(self) -> None:
        text = "realm-name: EXAMPLE.COM\n  domain-name: example.com\n"
        self.assertEqual(domain_from_realm(text), "example.com")
        self.assertEqual(domain_from_sssd("[sssd]\ndomains = lab.local\n"), "lab.local")
        self.assertTrue(user_is_domain_admin("administrador domain admins@example.com sudo"))
        self.assertFalse(user_is_domain_admin("administrador sudo adm"))

    def test_edid_nome(self) -> None:
        data = bytearray(128)
        data[54:57] = b"\x00\x00\x00"
        data[57] = 0xFC
        data[59:70] = b"DELL P2419H"
        self.assertEqual(edid_descriptor(bytes(data), 0xFC), "DELL P2419H")
        data[12:16] = (727613).to_bytes(4, "little")
        self.assertEqual(edid_numeric_serial(bytes(data)), "727613")
        line = (
            "HDMI-A-0 connected primary 1920x1080+0+0 "
            "(normal left inverted right x axis y axis) 527mm x 296mm"
        )
        monitors = parse_xrandr_monitors(line)
        self.assertEqual(monitors[0].resolution, "1920x1080")
        self.assertEqual(monitors[0].physical_size, "527mm x 296mm")
        rates = parse_xrandr_refresh(line + "\n   1920x1080    100.00*+  60.00\n")
        self.assertEqual(rates["HDMI-A-0"], "100 Hz")

    def test_hotspot_e_bluetooth(self) -> None:
        texto = "Supported interface modes:\n\t\t * managed\n\t\t * AP\n\t\t * AP/VLAN\nBand 1:\n"
        self.assertTrue(ap_mode_supported(texto))
        self.assertFalse(ap_mode_supported("Supported interface modes:\n\t\t * managed\n"))
        self.assertTrue(usb_has_bluetooth("e0", "01", "01"))
        self.assertFalse(usb_has_bluetooth("0e", "01", "00"))
        self.assertEqual(gpu_instalacao("AMD/ATI Phoenix1", "Ryzen 7 8700G w/ Radeon 780M Graphics"), "sim")
        self.assertEqual(gpu_instalacao("NVIDIA GeForce RTX 3060", "Ryzen 7 8700G"), "nao")
        self.assertEqual(cpu_place_and_socket("AM5", "Socket AM5"), ("Unico processador da placa", "AM5"))
        self.assertEqual(cpu_place_and_socket("CPU 1", "Socket AM5"), ("CPU 1", "AM5"))
        self.assertEqual(chipset_from_name("LOG-A620-DDR5", ""), "A620")
        self.assertTrue(slot_is_free("Available"))
        self.assertFalse(slot_is_free("In Use"))
        self.assertIn("video", slot_purpose("PCIEX16", "x16 PCI Express", "16x"))
        vazio = dict(_radio_rows(SimpleNamespace(wifi_adapter="", hotspot="", bluetooth="", hotspot_detail="", camera_integrada="", caixa_som="", saidas_audio="", slots=[])))
        self.assertEqual(vazio["Adaptador Wi-Fi"], "Sim, por USB")
        self.assertEqual(vazio["Adaptador Bluetooth"], "Sim, por USB")
        self.assertTrue(vazio["Hotspot"].startswith("Nao"))
        pronto = dict(_radio_rows(SimpleNamespace(wifi_adapter="phy0", hotspot="Sim", bluetooth="Sim", hotspot_detail="", camera_integrada="", caixa_som="", saidas_audio="Fone frontal, Line-out traseiro, HDMI", slots=[])))
        self.assertNotIn("Audio so no fone", pronto)
        self.assertNotIn("Caixa de som externa", pronto)
        self.assertEqual(pronto["Saidas de audio"], "Fone frontal, Line-out traseiro, HDMI")
        codec = "Pin Default 0x01014010: [Jack] Line Out at Ext Rear\nPin Default 0x02214030: [Jack] HP Out at Ext Front\nPin Default 0x411111f0: [N/A] Speaker at Ext Rear\n"
        self.assertEqual(saidas_from_codec(codec), ["Line-out traseiro", "Fone frontal"])
        self.assertTrue(hdmi_audio_active("monitor_present\t\t1\neld_valid\t\t1\n"))
        self.assertFalse(hdmi_audio_active("monitor_present\t\t0\neld_valid\t\t0\n"))
        self.assertEqual(nvme_link_label("8.0 GT/s PCIe", "4"), "PCIe 3.0 x4 (8 GT/s)")
        self.assertEqual(sata_speed_from_caps("ef36ff25"), "6 Gbps")
        self.assertTrue(sata_port_is_real("4004016"))
        self.assertFalse(sata_port_is_real("0"))
        partes = (
            "major minor  #blocks  name\n"
            " 259 0 500107608 nvme0n1\n"
            " 259 0 524288 nvme0n1p1\n"
            " 259 0 116695040 nvme0n1p2\n"
        )
        gaps = dict(unallocated_from_partitions(partes))
        self.assertGreater(gaps["nvme0n1"], 300 * 1024 * 1024 * 1024)
        self.assertEqual(pronto["Adaptador Wi-Fi"], "Sim. Ja ha adaptador.")
        self.assertEqual(pronto["Adaptador Bluetooth"], "Sim. Ja ha adaptador.")
        self.assertTrue(pronto["Hotspot"].startswith("Sim"))

    def test_edid_fabricacao_codigo_e_nativa(self) -> None:
        data = bytearray(128)
        data[10:12] = (0x0E1C).to_bytes(2, "little")
        data[16] = 12
        data[17] = 34
        data[54] = 1
        data[55] = 1
        data[56] = 128
        data[58] = 0x70
        data[59] = 56
        data[61] = 0x40
        self.assertEqual(edid_product_code(bytes(data)), "0E1C")
        self.assertEqual(edid_manufacture(bytes(data)), "2024 (semana 12)")
        self.assertEqual(edid_native_resolution(bytes(data)), "1920x1080")

    def test_volumes_dominio_e_edicao(self) -> None:
        device = {
            "children": [
                {"type": "part", "name": "nvme0n1p2", "fstype": "ext4", "label": "", "mountpoint": "/"},
            ]
        }
        self.assertEqual(volume_summary(device), "nvme0n1p2 ext4 /")
        sem_tipo = {"children": [{"type": "part", "name": "nvme0n1p2", "fstype": None, "mountpoint": "/"}]}
        self.assertEqual(volume_summary(sem_tipo, {"nvme0n1p2": ("ext4", "/")}), "nvme0n1p2 ext4 /")
        realm = "realm-name: EXAMPLE.COM\n  domain-name: example.com\n"
        self.assertEqual(domain_group_label(realm, "", ""), "Dominio: example.com")
        self.assertEqual(domain_group_label("", "", "WORKGROUP"), "Grupo: WORKGROUP")
        self.assertEqual(domain_group_label("", "localdomain", ""), "Nao ingressado")
        self.assertEqual(
            os_edition_version({"VERSION": "22.2 (Zena)", "VERSION_ID": "22.2", "VERSION_CODENAME": "zena"}),
            ("Zena", "22.2"),
        )


class ReportTest(unittest.TestCase):
    def test_fluxos_de_tombamento(self) -> None:
        hardware = _hardware()
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            with_number = write_inventory(
                FormData(patrimonio="123456", responsavel="Maria Silva", local="Sala 12", status="ativo"),
                hardware,
                "individual",
                folder,
            )
            without = write_inventory(
                FormData(responsavel="Maria Silva", local="Sala 12", chave_licenca="AAAAA-BBBBB"),
                hardware,
                "individual",
                folder,
            )
            self.assertEqual(with_number.spreadsheet, "123456_PC-LAB_inventario.xlsx")
            self.assertEqual(without.spreadsheet, "PC-LAB_inventario.xlsx")
            self.assertTrue(with_number.pdf.endswith("_relatorio_hardware.pdf"))

            book = load_workbook(with_number.spreadsheet_path)
            self.assertEqual(
                book.sheetnames,
                ["Capa", "Maquina", "Componentes", "Monitores", "Auditoria", "Expansao"],
            )
            capa_b = [cell.value for cell in book["Capa"]["B"]]
            self.assertIn("123456", capa_b)
            self.assertNotIn("FALTA numero de patrimonio", capa_b)
            self.assertEqual(book["Componentes"]["A1"].value, "Categoria")
            self.assertEqual(book["Componentes"]["D1"].value, "Nº serie")
            maquina_b = [cell.value for cell in book["Maquina"]["B"]]
            self.assertIn("123456", maquina_b)
            self.assertIn("Maria Silva", maquina_b)

            provisional = load_workbook(without.spreadsheet_path)
            capa_alerts = [cell.value for cell in provisional["Capa"]["B"]]
            self.assertIn("FALTA numero de patrimonio", capa_alerts)
            self.assertIn("AAAAA-BBBBB", [cell.value for cell in provisional["Maquina"]["B"]])
            audit = [cell.value for cell in provisional["Auditoria"]["A"]]
            self.assertIn("motherboard", audit)
            self.assertTrue(Path(without.pdf_path).read_bytes().startswith(b"%PDF"))
            self.assertTrue(Path(with_number.pdf_path).read_bytes().startswith(b"%PDF"))

    def test_auditoria_sem_serie_da_placa(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            result = write_inventory(
                FormData(patrimonio="1", responsavel="A", local="B"),
                _hardware(),
                "individual",
                Path(tmp),
            )
            book = load_workbook(result.spreadsheet_path)
            audit = [cell.value for cell in book["Auditoria"]["A"]]
            self.assertIn("motherboard", audit)
            capa = [cell.value for cell in book["Capa"]["B"]]
            self.assertNotIn("FALTA numero de patrimonio", capa)

    def test_chave_do_firmware_entra_no_relatorio(self) -> None:
        hardware = _hardware()
        hardware.license_key = "AAAAA-BBBBB-CCCCC-DDDDD-EEEEE"
        hardware.license_key_partial = partial_key(hardware.license_key)
        with tempfile.TemporaryDirectory() as tmp:
            result = write_inventory(
                FormData(patrimonio="1", responsavel="A", local="B"),
                hardware,
                "individual",
                Path(tmp),
            )
            book = load_workbook(result.spreadsheet_path)
            capa = [cell.value for cell in book["Capa"]["B"]]
            maquina = [cell.value for cell in book["Maquina"]["B"]]
            self.assertIn(hardware.license_key, capa)
            self.assertIn(hardware.license_key, maquina)
            self.assertIn("EEEEE", maquina)
            self.assertIn("Chave do Linux", [cell.value for cell in book["Maquina"]["A"]])
            self.assertNotIn("Chave (Registro Windows)", [cell.value for cell in book["Maquina"]["A"]])

    def test_msdm(self) -> None:
        data = bytearray(56)
        data[:4] = b"MSDM"
        key = b"AAAAA-BBBBB-CCCCC-DDDDD-EEEEE"
        data[52:56] = len(key).to_bytes(4, "little")
        data += key
        self.assertEqual(key_from_msdm(bytes(data)), key.decode())

    def test_licenca_do_linux(self) -> None:
        text = "Files: *\nLicense: GPL-2.0+\n"
        self.assertEqual(license_from_copyright(text), "GPL-2.0+")
        self.assertTrue(is_free_license("GPL-2.0+"))
        self.assertFalse(is_free_license("AAAAA-BBBBB-CCCCC-DDDDD-EEEEE"))
        self.assertEqual(partial_key("GPL-2.0+"), "GPL-2.0+")
        self.assertEqual(token_from_ubuntu_pro({"contractToken": "C1a2b3c4d5e6"}), "C1a2b3c4d5e6")
        self.assertEqual(token_from_ubuntu_pro({"token": "a.b.c"}), "")


class LiveComputerTest(unittest.TestCase):
    def test_coleta_deste_computador(self) -> None:
        hardware = collect_hardware()
        self.assertTrue(hardware.hostname)
        self.assertTrue(hardware.os_name)
        self.assertTrue(hardware.os_version)
        self.assertTrue(hardware.license_key)
        self.assertTrue(is_free_license(hardware.license_key))
        self.assertTrue(hardware.domain_group)
        self.assertTrue(hardware.cpu_model)
        self.assertTrue(hardware.cpu_family)
        self.assertIn("bits", hardware.cpu_address)
        self.assertEqual(hardware.hotspot, "Nao")
        self.assertEqual(hardware.bluetooth, "Nao")
        self.assertEqual(hardware.camera_integrada, "Nao")
        self.assertEqual(hardware.caixa_som, "Nao")
        self.assertIn("Fone frontal", hardware.saidas_audio)
        self.assertIn("Line-out traseiro", hardware.saidas_audio)
        self.assertIn("HDMI", hardware.saidas_audio)
        self.assertEqual(hardware.ssd_speed, "PCIe 3.0 x4 (8 GT/s)")
        self.assertEqual(hardware.sata_ports_total, "4")
        self.assertEqual(hardware.sata_ports_free, "4")
        self.assertEqual(hardware.sata_speed, "6 Gbps")
        self.assertTrue(hardware.cpu_l3 or hardware.cpu_l2)
        self.assertTrue(hardware.disks)
        self.assertTrue(any(disk.volumes for disk in hardware.disks))
        self.assertTrue(hardware.gpus)
        self.assertTrue(hardware.gpus[0].name)
        self.assertTrue(hardware.memory_total)
        with tempfile.TemporaryDirectory() as tmp:
            result = write_inventory(
                FormData(patrimonio="999", responsavel="Responsável local", local="Computador pessoal"),
                hardware,
                "individual",
                Path(tmp),
            )
            book = load_workbook(result.spreadsheet_path)
            values = [cell.value for cell in book["Maquina"]["B"]]
            self.assertIn(hardware.hostname, values)
            self.assertIn("Responsável local", values)
            self.assertIn(hardware.os_version, values)
            self.assertIn(hardware.license_key, values)
            self.assertIn("Chave do Linux", [cell.value for cell in book["Maquina"]["A"]])
            maquina_a = [cell.value for cell in book["Maquina"]["A"]]
            self.assertIn("Hotspot", maquina_a)
            self.assertIn("Adaptador Wi-Fi", maquina_a)
            self.assertIn("Adaptador Bluetooth", maquina_a)
            self.assertIn("Camera integrada", maquina_a)
            self.assertNotIn("Caixa de som externa", maquina_a)
            self.assertNotIn("Audio so no fone", maquina_a)
            self.assertIn("Saidas de audio", maquina_a)
            self.assertIn("Slots de RAM vazios", maquina_a)
            self.assertIn("Slots PCI vazios", maquina_a)
            self.assertIn("Nucleos", maquina_a)
            self.assertIn("Threads", maquina_a)
            self.assertIn("Soquetes totais", maquina_a)
            self.assertIn("HDD que ainda cabem", maquina_a)
            self.assertIn("SSD que ainda cabem", maquina_a)
            self.assertIn("Capacidade maxima", maquina_a)
            self.assertIn("Onboard", maquina_a)
            self.assertNotIn("Pode ser hotspot?", maquina_a)
            self.assertNotIn("Camera integrada?", maquina_a)
            radio = dict(_radio_rows(hardware))
            self.assertEqual(radio["Adaptador Wi-Fi"], "Sim, por USB")
            self.assertEqual(radio["Adaptador Bluetooth"], "Sim, por USB")
            self.assertTrue(radio["Hotspot"].startswith("Nao"))
            self.assertNotIn("Audio so no fone", radio)
            self.assertNotIn("Caixa de som externa", radio)
            self.assertFalse(any(value and str(value).startswith("Hoje:") for value in book["Maquina"]["B"]))
            maquina_a = [cell.value for cell in book["Maquina"]["A"]]
            self.assertNotIn("5. RESPOSTAS DO EQUIPAMENTO", maquina_a)
            self.assertNotIn("Existe placa de video?", maquina_a)
            expansao_a = [cell.value for cell in book["Expansao"]["A"]]
            expansao_b = [cell.value for cell in book["Expansao"]["B"]]
            self.assertIn("Velocidade do SSD", expansao_a)
            self.assertIn("Controlador SATA", expansao_a)
            self.assertIn("Portas SATA livres", expansao_a)
            self.assertIn("PCIe 3.0 x4 (8 GT/s)", expansao_b)
            self.assertFalse(any(value and "Windows" in str(value) for value in expansao_a + expansao_b))
            self.assertTrue(any(value and "dmidecode" in str(value) for value in expansao_b))
            self.assertIn("O que da para melhorar", expansao_a)
            self.assertTrue(hardware.disk_space)
            self.assertIn("sem particao", hardware.disk_space)
            self.assertIn("4) PROCESSADOR", expansao_a)
            self.assertIn("5) REDE E PERIFERICOS", expansao_a)
            self.assertTrue(any(value and "4 de 4" in str(value) for value in expansao_b))
            notas = [cell.value for cell in book["Componentes"]["R"]]
            self.assertTrue(any(value and "Instalacao: Onboard, no processador" in str(value) for value in notas))
            self.assertTrue(any(value and "Outra placa interna" in str(value) for value in notas))
            slots = [cell.value for cell in book["Componentes"]["Q"]]
            self.assertTrue(any(value and "slots vazios nao lidos" in str(value) for value in slots))
            self.assertTrue(any(value and "SATA da placa, 6 Gbps, 4 portas" in str(value) for value in slots))
            self.assertNotIn("Registro Windows", Path(result.pdf_path).read_bytes().decode("latin-1", "ignore"))
            self.assertIn(hardware.domain_group, values)
            componentes = [cell.value for cell in book["Componentes"]["R"]]
            self.assertTrue(any(value and "Family=" in str(value) for value in componentes))
            self.assertEqual(
                book.sheetnames,
                ["Capa", "Maquina", "Componentes", "Monitores", "Auditoria", "Expansao"],
            )
            self.assertGreater(Path(result.pdf_path).stat().st_size, 1000)


if __name__ == "__main__":
    unittest.main()
